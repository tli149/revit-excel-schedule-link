"""Generate testdata/Demo_Schedules.xlsx, an English sample workbook for
documentation screenshots and manual checks.

It exercises the features shown in the user docs: a merged title row, plain
text, whole numbers, thousands separators, percentages, unit literals in
custom formats (A, kVA) and a fraction format.

Usage: python testdata/make_demo_workbook.py   (requires openpyxl)
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Demo_Schedules.xlsx")


def door_schedule(ws):
    ws.title = "Door Schedule"
    ws.merge_cells("A1:G1")
    ws["A1"] = "Door Schedule - Level 1"
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    headers = ["Mark", "Room", "Width (mm)", "Height (mm)", "Fire Rating",
               "Hardware Set", "Remarks"]
    ws.append(headers)
    rows = [
        ("D101", "Lobby", 1800, 2400, "-", "HW-01", "Double leaf, glazed"),
        ("D102", "Reception", 900, 2100, "-", "HW-02", ""),
        ("D103", "Office 1", 900, 2100, "-", "HW-02", ""),
        ("D104", "Office 2", 900, 2100, "-", "HW-02", ""),
        ("D105", "Meeting Room", 1000, 2100, "-", "HW-03", "Acoustic seal"),
        ("D106", "Stair A", 1000, 2100, "60 min", "HW-04", "Self-closing"),
        ("D107", "Electrical Room", 1000, 2100, "60 min", "HW-05", "Lockable"),
        ("D108", "Storage", 900, 2100, "30 min", "HW-05", ""),
    ]
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=3, min_col=3, max_col=4):
        for cell in row:
            cell.number_format = "#,##0"
    widths = [8, 18, 12, 12, 12, 14, 22]
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = width


def panel_loads(wb):
    ws = wb.create_sheet("Panel Loads")
    ws.merge_cells("A1:F1")
    ws["A1"] = "Panel LP-1 Load Summary"
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["Circuit", "Description", "Current", "Load", "Demand", "Conduit (in)"])
    rows = [
        (1, "Lighting - Lobby", 6.4, 1.47, 1.0, 0.5),
        (2, "Lighting - Offices", 8.2, 1.89, 1.0, 0.5),
        (3, "Receptacles - Offices", 12.5, 2.88, 0.5, 0.75),
        (4, "Receptacles - Meeting", 9.1, 2.09, 0.5, 0.75),
        (5, "Fan Coil Unit FCU-1", 4.8, 1.10, 0.8, 0.5),
        (6, "Water Heater WH-1", 18.0, 4.14, 1.0, 1.25),
    ]
    for row in rows:
        ws.append(row)
    for (current, load, demand, conduit) in ws.iter_rows(min_row=3, min_col=3, max_col=6):
        current.number_format = '0.0 "A"'
        load.number_format = '0.00 "kVA"'
        demand.number_format = "0%"
        conduit.number_format = "# ?/?"
    widths = [9, 26, 10, 12, 10, 13]
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = width


def main():
    wb = Workbook()
    door_schedule(wb.active)
    panel_loads(wb)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
