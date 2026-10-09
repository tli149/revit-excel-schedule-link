# -*- coding: utf-8 -*-
"""Generate Step0_TestBook.xlsx.

The workbook covers core CellDisplayFormatter formats (integers, decimals,
grouping, percentages, dates, times, leading zeros, and cached formulas),
merged cells, and multiple sheets. Its Japanese sheet names and cell values are
intentional CJK test data and must remain unchanged.

Formula cache values:
  openpyxl does not write a calculated result (<v> element) for formulas, so
  this script injects <v> directly into the sheet XML inside the XLSX archive
  after saving. The Open XML SDK reads this cached value; without it a formula
  cell appears empty.
"""
import re
import zipfile
import shutil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / "Step0_TestBook.xlsx"

FONT = "Arial"
FORMULA_VALUE = 1234.567 * 2  # Cached value for formula =B5*2 in B13.

def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Step0"

    ws["A1"] = "ESL Step0 テストブック(書式ケース)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    ws.merge_cells("A1:C1")

    headers = ["ケース", "値(書式適用)", "期待表示"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(name=FONT, bold=True)

    # (case name, value, number_format, expected display string)
    rows = [
        ("整数 0",            1234.567,  "0",          "1235"),
        ("小数 0.00",         1234.567,  "0.00",       "1234.57"),
        ("桁区切り #,##0",     1234567.89, "#,##0",      "1,234,568"),
        ("桁区切り #,##0.00",  1234567.89, "#,##0.00",   "1,234,567.89"),
        ("パーセント 0%",      0.1235,    "0%",         "12%"),
        ("パーセント 0.00%",   0.1235,    "0.00%",      "12.35%"),
        ("日付 yyyy/m/d",     None,      "yyyy/m/d",   "2026/6/10"),
        ("時刻 h:mm",         None,      "h:mm",       "13:45"),
        ("先頭ゼロ 00000",     123,       "00000",      "00123"),
        ("数式 =B5*2",        None,      "#,##0.00",   "2,469.13"),
        ("文字列(inlineStr)", "日本語テキスト ABC123", "@", "日本語テキスト ABC123"),
        ("ブール",            True,      None,         "TRUE"),
    ]
    import datetime
    r = 4
    formula_row = None
    for name, value, fmt, expected in rows:
        ws.cell(row=r, column=1, value=name).font = Font(name=FONT)
        c = ws.cell(row=r, column=2)
        if name.startswith("日付"):
            c.value = datetime.date(2026, 6, 10)
        elif name.startswith("時刻"):
            c.value = datetime.time(13, 45, 0)
        elif name.startswith("数式"):
            c.value = "=B5*2"
            formula_row = r
        else:
            c.value = value
        if fmt:
            c.number_format = fmt
        c.font = Font(name=FONT)
        e = ws.cell(row=r, column=3, value=expected)
        e.font = Font(name=FONT)
        e.number_format = "@"
        r += 1

    # Merged cell with a value.
    ws.cell(row=r + 1, column=1, value="結合セル A%d:C%d" % (r + 1, r + 1)).font = Font(
        name=FONT, italic=True
    )
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=3)
    ws.cell(row=r + 1, column=1).alignment = Alignment(horizontal="center")

    for col, w in (("A", 24), ("B", 18), ("C", 22)):
        ws.column_dimensions[col].width = w

    # Second sheet. The sheet name and A1 value are intentional CJK test data.
    ws2 = wb.create_sheet("データ")
    ws2["A1"] = "日本語シートの A1"
    ws2["A1"].font = Font(name=FONT)
    for i in range(2, 5):
        for j in range(1, 4):
            cc = ws2.cell(row=i, column=j, value=(i - 1) * 10 + j)
            cc.font = Font(name=FONT)
            cc.number_format = "#,##0"
    ws2.column_dimensions["A"].width = 20

    wb.save(OUT)
    return formula_row


def inject_formula_cached_value(formula_row):
    """Inject a cached <v> value into sheet1 formula cell B{row}."""
    src = OUT
    tmp = OUT.with_suffix(".tmp.xlsx")
    cell_ref = "B%d" % formula_row
    # openpyxl writes <f>...</f><v /> with an empty cache; replace that <v />.
    pat = re.compile(
        r'(<c r="%s"[^>]*>)(<f>[^<]*</f>)<v\s*/>(</c>)' % cell_ref
    )
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(
        tmp, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                new_text, n = pat.subn(
                    r"\1\2<v>%s</v>\3" % repr(FORMULA_VALUE), text
                )
                if n != 1:
                    raise RuntimeError(
                        "Formula cell %s was not found (replacement count=%d)" % (cell_ref, n)
                    )
                data = new_text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(src))


def verify():
    """Verify the formula cache, merges, formats, and sheet names."""
    from openpyxl import load_workbook

    wb_v = load_workbook(OUT, data_only=True)   # Read cached values.
    wb_f = load_workbook(OUT, data_only=False)  # Read formulas.
    ws_v, ws_f = wb_v["Step0"], wb_f["Step0"]
    # The Japanese sheet name is intentional CJK test data.
    assert wb_v.sheetnames == ["Step0", "データ"], wb_v.sheetnames

    found = None
    for row in ws_f.iter_rows(min_col=2, max_col=2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                found = c
    assert found is not None, "No formula cell was found"
    cached = ws_v[found.coordinate].value
    assert cached is not None and abs(cached - FORMULA_VALUE) < 1e-9, (
        "Cached value mismatch: %r" % cached
    )
    assert len(ws_f.merged_cells.ranges) == 2, ws_f.merged_cells.ranges
    print("verify OK: formula %s cache=%s / merges=%s / sheets=%s" % (
        found.coordinate, cached, list(ws_f.merged_cells.ranges), wb_v.sheetnames
    ))


if __name__ == "__main__":
    formula_row = build()
    inject_formula_cached_value(formula_row)
    verify()
    print("Generated:", OUT)
