# -*- coding: utf-8 -*-
"""Step5_FormatsBook.xlsx generator.

A small demo workbook for the Step 5 number-format work: unit-suffix custom
formats and fraction formats. Column B holds the value with the format applied;
column C is the expected display string. Used for manual Revit verification of
v0.6.0 (the unit tests already cover the formatter logic).
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).parent / "Step5_FormatsBook.xlsx"
FONT = "Arial"


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Formats"

    ws["A1"] = "ESL Step 5 number-format demo"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)

    headers = ["Case", "Value (formatted)", "Expected"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(name=FONT, bold=True)

    rows = [
        ("Unit A",        12.345, '0.0 "A"',    "12.3 A"),
        ("Unit kVA",      12.345, '0.00 "kVA"', "12.35 kVA"),
        ("Unit CFM",      1234,   '0 "CFM"',    "1234 CFM"),
        ("Unit GPM",      12.3,   '0.0 "GPM"',  "12.3 GPM"),
        ("Unit negative", -5.0,   '0.0 "A"',    "-5.0 A"),
        ("Currency $",    1234.56, '"$"#,##0.00', "$1,234.56"),
        ("Fraction 1/2",  1.5,    "# ?/?",      "1 1/2"),
        ("Fraction 1/4",  2.25,   "# ?/?",      "2 1/4"),
        ("Fraction 7/10", 2.7,    "# ??/??",    "2 7/10"),
        ("Fraction 1/8",  3.125,  "# ??/??",    "3 1/8"),
        ("Fraction whole", 2.0,   "# ?/?",      "2"),
    ]
    r = 4
    for name, value, fmt, expected in rows:
        ws.cell(row=r, column=1, value=name).font = Font(name=FONT)
        c = ws.cell(row=r, column=2, value=value)
        c.number_format = fmt
        c.font = Font(name=FONT)
        e = ws.cell(row=r, column=3, value=expected)
        e.number_format = "@"
        e.font = Font(name=FONT)
        r += 1

    for col, w in (("A", 18), ("B", 18), ("C", 14)):
        ws.column_dimensions[col].width = w

    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
